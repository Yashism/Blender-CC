"""Put the slide recordings into a ready-to-use PowerPoint file.

Each MP4 fills one 16:9 slide and is set to Start: Automatically and Loop until
Stopped, so it plays as soon as the slide comes up. Copy the slide into any deck.

    python make_pptx.py --out ../out/PVR_Entry_Counter.pptx \
        --video ../out/PVR_Entry_Counter_AI_Footage.mp4 --video ../out/PVR_Entry_Counter_Pure_Code.mp4
"""
import argparse
import subprocess
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

P = "http://schemas.openxmlformats.org/presentationml/2006/main"

# PowerPoint's own timing tree for "Start: Automatically" + "Loop until Stopped"
TIMING = """<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{dur}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="80000"><p:cTn id="7" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn>
<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""


def duration_ms(path):
    out = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(path)],
                         capture_output=True, text=True).stdout
    return int(float(out) * 1000)


def poster(path, at, out):
    subprocess.run(["ffmpeg", "-v", "error", "-y", "-ss", str(at), "-i", str(path), "-frames:v", "1", "-q:v", "2", str(out)], check=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--video", action="append", required=True)
    ap.add_argument("--poster-at", type=float, default=12.0)
    args = ap.parse_args()

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in, 16:9
    blank = prs.slide_layouts[6]
    tmp = Path(tempfile.mkdtemp())
    for i, v in enumerate(args.video):
        v = Path(v)
        slide = prs.slides.add_slide(blank)
        jpg = tmp / f"poster_{i}.jpg"
        poster(v, args.poster_at, jpg)
        movie = slide.shapes.add_movie(str(v), 0, 0, prs.slide_width, prs.slide_height,
                                       poster_frame_image=str(jpg), mime_type="video/mp4")
        movie.name = v.stem
        sld = slide._element
        old = sld.find(f"{{{P}}}timing")
        if old is not None:
            sld.remove(old)
        timing = etree.fromstring(TIMING.format(P=P, dur=duration_ms(v), spid=movie.shape_id))
        # p:timing goes after p:clrMapOvr / before p:extLst
        ext = sld.find(f"{{{P}}}extLst")
        if ext is not None:
            ext.addprevious(timing)
        else:
            sld.append(timing)
    prs.save(args.out)
    print(f"wrote {args.out} ({len(args.video)} slides)")


if __name__ == "__main__":
    main()
